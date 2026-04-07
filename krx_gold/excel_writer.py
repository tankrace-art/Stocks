"""KRX 금시장 데이터 엑셀 저장 + 차트 생성

기능:
- 누적 저장 (한 파일에 계속 추가, 일자 중복 제거)
- 사용자 양식: 일자/종가/대비/등락률/시가/고가/저가/거래량/거래대금
- 가격/거래량 그래프 자동 생성 (Excel에 임베드)
"""

import os
from io import BytesIO
from datetime import datetime

import pandas as pd
from openpyxl import load_workbook, Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils.dataframe import dataframe_to_rows

import matplotlib
matplotlib.use("Agg")  # GUI 백엔드 비활성화
import matplotlib.pyplot as plt
import matplotlib.dates as mdates

# ── 한글 폰트 설정 ──
import platform
if platform.system() == "Windows":
    plt.rcParams["font.family"] = "Malgun Gothic"
elif platform.system() == "Darwin":
    plt.rcParams["font.family"] = "AppleGothic"
else:
    plt.rcParams["font.family"] = "DejaVu Sans"
plt.rcParams["axes.unicode_minus"] = False

# 사용자 요청 컬럼 순서
USER_COLUMNS = ["일자", "종가", "대비", "등락률", "시가", "고가", "저가", "거래량", "거래대금"]

# KRX API 응답 → 사용자 컬럼 매핑
API_COL_MAP = {
    "BAS_DD": "일자",
    "TRD_DD": "일자",
    "TDD_CLSPRC": "종가",
    "CMPPREVDD_PRC": "대비",
    "FLUC_RT": "등락률",
    "TDD_OPNPRC": "시가",
    "TDD_HGPRC": "고가",
    "TDD_LWPRC": "저가",
    "ACC_TRDVOL": "거래량",
    "ACC_TRDVAL": "거래대금",
    "ISU_NM": "종목명",
    "ISU_CD": "종목코드",
}


def normalize_dataframe(df: pd.DataFrame) -> pd.DataFrame:
    """KRX API 응답을 사용자 양식으로 변환"""
    if df.empty:
        return df

    # 컬럼명 매핑
    rename = {old: new for old, new in API_COL_MAP.items() if old in df.columns}
    df = df.rename(columns=rename)

    # 일자 형식 통일 (YYYY/MM/DD)
    if "일자" in df.columns:
        df["일자"] = df["일자"].astype(str).str.replace("-", "/").str.replace(".", "/")
        # YYYYMMDD 형식이면 슬래시 추가
        df["일자"] = df["일자"].apply(_format_date)

    # 숫자 컬럼 정리 (콤마 제거 후 숫자 변환)
    num_cols = ["종가", "대비", "등락률", "시가", "고가", "저가", "거래량", "거래대금"]
    for col in num_cols:
        if col in df.columns:
            df[col] = pd.to_numeric(
                df[col].astype(str).str.replace(",", "", regex=False),
                errors="coerce",
            )

    # 사용자 요청 컬럼 순서로 정렬 (없는 컬럼은 제외)
    cols = [c for c in USER_COLUMNS if c in df.columns]
    other_cols = [c for c in df.columns if c not in cols]
    df = df[cols + other_cols]

    return df


def _format_date(val: str) -> str:
    """다양한 일자 형식을 YYYY/MM/DD로 통일"""
    s = str(val).strip().replace("-", "").replace("/", "").replace(".", "")
    if len(s) == 8 and s.isdigit():
        return f"{s[:4]}/{s[4:6]}/{s[6:]}"
    return val


def save_cumulative_excel(
    df_new: pd.DataFrame,
    filepath: str,
    sheet_name: str = "금시세",
    item_name: str = None,
    log_fn=None,
) -> int:
    """엑셀 파일에 누적 저장 (일자 중복 제거 + 그래프 생성)

    Parameters
    ----------
    df_new : pd.DataFrame
        새로 추가할 데이터
    filepath : str
        엑셀 파일 경로
    sheet_name : str
        시트 이름
    item_name : str, optional
        종목명 (기간별 시세일 때만 의미 있음)
    log_fn : callable

    Returns
    -------
    int
        최종 저장된 행 수
    """
    if df_new.empty:
        return 0

    df_new = normalize_dataframe(df_new)

    # 종목 필터: 종목별 그래프를 위해 단일 종목만 처리
    if item_name and "종목명" in df_new.columns:
        df_new = df_new[df_new["종목명"].str.contains(item_name, na=False)]

    # ── 기존 파일이 있으면 병합 ──
    if os.path.exists(filepath):
        try:
            df_old = pd.read_excel(filepath, sheet_name=sheet_name)
            if log_fn:
                log_fn(f"  기존 파일에서 {len(df_old)}행 로드")

            # 컬럼 합집합으로 정렬
            all_cols = list(dict.fromkeys(list(df_old.columns) + list(df_new.columns)))
            df_old = df_old.reindex(columns=all_cols)
            df_new = df_new.reindex(columns=all_cols)

            df = pd.concat([df_old, df_new], ignore_index=True)
        except Exception as e:
            if log_fn:
                log_fn(f"  기존 파일 읽기 실패, 새로 생성: {e}")
            df = df_new
    else:
        df = df_new

    # ── 일자 기준 중복 제거 (최신 우선) ──
    if "일자" in df.columns:
        df = df.drop_duplicates(subset=["일자"], keep="last")
        df = df.sort_values("일자").reset_index(drop=True)

    if log_fn:
        log_fn(f"  최종 저장 {len(df)}행 (누적)")

    # ── 엑셀 작성 ──
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name

    # 헤더 스타일
    header_fill = PatternFill("solid", fgColor="4472C4")
    header_font = Font(bold=True, color="FFFFFF", size=11)
    center = Alignment(horizontal="center", vertical="center")
    thin = Side(border_style="thin", color="CCCCCC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    # 데이터 쓰기
    for row in dataframe_to_rows(df, index=False, header=True):
        ws.append(row)

    # 헤더 스타일 적용
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center
        cell.border = border

    # 데이터 행 스타일
    num_format_int = "#,##0"
    num_format_pct = "0.00"
    for row_idx in range(2, ws.max_row + 1):
        for col_idx, col_name in enumerate(df.columns, start=1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.border = border
            if col_name in ("종가", "시가", "고가", "저가", "대비"):
                cell.number_format = num_format_int
                cell.alignment = Alignment(horizontal="right")
            elif col_name == "등락률":
                cell.number_format = num_format_pct
                cell.alignment = Alignment(horizontal="right")
            elif col_name in ("거래량", "거래대금"):
                cell.number_format = num_format_int
                cell.alignment = Alignment(horizontal="right")
            elif col_name == "일자":
                cell.alignment = Alignment(horizontal="center")

    # 컬럼 너비 자동
    width_map = {
        "일자": 12, "종가": 12, "대비": 10, "등락률": 10,
        "시가": 12, "고가": 12, "저가": 12,
        "거래량": 14, "거래대금": 18, "종목명": 20, "종목코드": 14,
    }
    for col_idx, col_name in enumerate(df.columns, start=1):
        col_letter = ws.cell(row=1, column=col_idx).column_letter
        ws.column_dimensions[col_letter].width = width_map.get(col_name, 12)

    ws.freeze_panes = "A2"

    # ── 차트 생성 (데이터가 2일 이상일 때만) ──
    if len(df) >= 2 and "일자" in df.columns and "종가" in df.columns:
        chart_img = create_chart(df, item_name)
        if chart_img:
            img = XLImage(chart_img)
            img.width = 720
            img.height = 400
            anchor_col = ws.max_column + 2
            anchor_letter = ws.cell(row=1, column=anchor_col).column_letter
            ws.add_image(img, f"{anchor_letter}1")

    wb.save(filepath)
    return len(df)


def create_chart(df: pd.DataFrame, title_suffix: str = None) -> BytesIO:
    """가격 + 거래량 그래프 생성 (PNG → BytesIO)"""
    if "일자" not in df.columns or "종가" not in df.columns:
        return None

    df = df.copy()
    df["_date"] = pd.to_datetime(
        df["일자"].astype(str).str.replace("/", "-"), errors="coerce"
    )
    df = df.dropna(subset=["_date"]).sort_values("_date")

    if len(df) < 2:
        return None

    fig, ax1 = plt.subplots(figsize=(12, 6), dpi=100)

    # 가격 라인
    ax1.plot(df["_date"], df["종가"], color="#1f77b4", linewidth=2, label="종가")
    ax1.set_xlabel("일자", fontsize=11)
    ax1.set_ylabel("가격 (원/g)", color="#1f77b4", fontsize=11)
    ax1.tick_params(axis="y", labelcolor="#1f77b4")
    ax1.grid(True, alpha=0.3)
    ax1.yaxis.set_major_formatter(plt.FuncFormatter(lambda x, _: f"{int(x):,}"))

    # 거래량 막대 (오른쪽 축)
    if "거래량" in df.columns:
        ax2 = ax1.twinx()
        ax2.bar(df["_date"], df["거래량"], color="#ff7f0e", alpha=0.4,
                width=0.8, label="거래량")
        ax2.set_ylabel("거래량 (g)", color="#ff7f0e", fontsize=11)
        ax2.tick_params(axis="y", labelcolor="#ff7f0e")
        ax2.yaxis.set_major_formatter(plt.FuncFormatter(lambda x, _: f"{int(x):,}"))

    # 제목
    title = "KRX 금시장 시세"
    if title_suffix:
        title += f" - {title_suffix}"
    start_str = df["_date"].iloc[0].strftime("%Y.%m.%d")
    end_str = df["_date"].iloc[-1].strftime("%Y.%m.%d")
    title += f" ({start_str} ~ {end_str})"
    plt.title(title, fontsize=13, fontweight="bold", pad=15)

    # x축 날짜 포맷
    if len(df) > 30:
        ax1.xaxis.set_major_locator(mdates.MonthLocator())
        ax1.xaxis.set_major_formatter(mdates.DateFormatter("%Y-%m"))
    else:
        ax1.xaxis.set_major_locator(mdates.AutoDateLocator())
        ax1.xaxis.set_major_formatter(mdates.DateFormatter("%m-%d"))
    plt.setp(ax1.xaxis.get_majorticklabels(), rotation=30, ha="right")

    plt.tight_layout()

    buf = BytesIO()
    plt.savefig(buf, format="png", dpi=100, bbox_inches="tight")
    plt.close(fig)
    buf.seek(0)
    return buf
