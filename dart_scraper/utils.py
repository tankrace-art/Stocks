"""유틸리티 함수 - 투자분석용 엑셀 출력"""

import os
import re

import pandas as pd
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side, numbers
from openpyxl.utils import get_column_letter


def format_amount(amount_str: str) -> str:
    """금액 문자열을 읽기 쉬운 형태로 변환 (억 원 단위)"""
    if not amount_str or amount_str == "-":
        return "-"
    try:
        amount = int(str(amount_str).replace(",", ""))
        if abs(amount) >= 100_000_000:
            return f"{amount / 100_000_000:,.1f}억"
        elif abs(amount) >= 10_000:
            return f"{amount / 10_000:,.0f}만"
        return f"{amount:,}"
    except (ValueError, AttributeError):
        return str(amount_str)


def parse_period(period_str: str) -> tuple[int, int, list[str]]:
    """기간 문자열 파싱

    지원 형식:
        "2023" → (2023, 2023, ["사업보고서"])
        "2020-2023" → (2020, 2023, ["사업보고서"])
        "2023 분기" → (2023, 2023, ["1분기", "반기", "3분기", "사업보고서"])
        "2020-2023 분기" → (2020, 2023, ["1분기", "반기", "3분기", "사업보고서"])
    """
    period_str = period_str.strip()
    quarterly = False

    if "분기" in period_str:
        quarterly = True
        period_str = period_str.replace("분기", "").strip()

    match = re.match(r"(\d{4})\s*[-~]\s*(\d{4})", period_str)
    if match:
        start_year = int(match.group(1))
        end_year = int(match.group(2))
    else:
        match = re.match(r"(\d{4})", period_str)
        if match:
            start_year = end_year = int(match.group(1))
        else:
            raise ValueError(
                f"기간 형식을 인식할 수 없습니다: '{period_str}'\n"
                "예시: '2023', '2020-2023', '2023 분기', '2020-2023 분기'"
            )

    if quarterly:
        report_types = ["1분기", "반기", "3분기", "사업보고서"]
    else:
        report_types = ["사업보고서"]

    return start_year, end_year, report_types


def save_investment_excel(
    filepath: str,
    company_name: str,
    company_info: dict,
    metrics_df: pd.DataFrame,
    cross_tables: dict[str, pd.DataFrame],
    business_data: dict[int, dict[str, list[pd.DataFrame]]],
    start_year: int,
    end_year: int,
) -> str:
    """투자분석용 엑셀 파일 저장

    시트 구성:
    1. 투자지표_요약 - 핵심 투자지표 (ROE, 부채비율, 영업이익률 등)
    2. 재무상태표 - 연도별 크로스 비교
    3. 손익계산서 - 연도별 크로스 비교
    4. 현금흐름표 - 연도별 크로스 비교
    5. 사업내용_YYYY - 연도별 사업보고서 핵심 데이터
    """
    os.makedirs(os.path.dirname(filepath) or ".", exist_ok=True)

    with pd.ExcelWriter(filepath, engine="openpyxl") as writer:
        # ── 시트 1: 투자지표 요약 ──
        if metrics_df is not None and not metrics_df.empty:
            metrics_df.to_excel(writer, sheet_name="투자지표_요약", index=False)

        # ── 시트 2~4: 재무제표 크로스 테이블 ──
        stmt_order = ["재무상태표", "손익계산서", "포괄손익계산서", "현금흐름표", "자본변동표"]
        for stmt_name in stmt_order:
            if stmt_name in cross_tables:
                df = cross_tables[stmt_name]
                if not df.empty:
                    safe_name = stmt_name[:31]
                    df.to_excel(writer, sheet_name=safe_name, index=False)

        # ── 시트 5+: 사업내용 연도별 ──
        if business_data:
            for year in sorted(business_data.keys()):
                sections = business_data[year]
                # 한 연도의 모든 섹션을 하나의 시트에 합침
                combined_rows = []
                for section_key, tables in sections.items():
                    from .report_parser import INVESTMENT_SECTIONS
                    section_desc = INVESTMENT_SECTIONS.get(
                        section_key, {}
                    ).get("description", section_key)

                    # 섹션 헤더 행 추가
                    combined_rows.append(
                        pd.DataFrame([{
                            "구분": f"▶ {section_desc}",
                        }])
                    )

                    for df in tables:
                        combined_rows.append(df)
                        # 빈 행 구분
                        combined_rows.append(pd.DataFrame([{}]))

                if combined_rows:
                    combined_df = pd.concat(combined_rows, ignore_index=True)
                    sheet_name = f"사업내용_{year}"[:31]
                    combined_df.to_excel(writer, sheet_name=sheet_name, index=False)

        # ── 엑셀 스타일링 ──
        wb = writer.book
        _style_workbook(wb, company_name, start_year, end_year)

    return filepath


def _style_workbook(wb, company_name: str, start_year: int, end_year: int):
    """엑셀 워크북 스타일링"""
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_font = Font(name="맑은 고딕", bold=True, color="FFFFFF", size=11)
    section_fill = PatternFill(start_color="D6E4F0", end_color="D6E4F0", fill_type="solid")
    section_font = Font(name="맑은 고딕", bold=True, size=11)
    data_font = Font(name="맑은 고딕", size=10)
    thin_border = Border(
        left=Side(style="thin", color="D0D0D0"),
        right=Side(style="thin", color="D0D0D0"),
        top=Side(style="thin", color="D0D0D0"),
        bottom=Side(style="thin", color="D0D0D0"),
    )

    for ws in wb.worksheets:
        # 헤더 행 스타일
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = thin_border

        # 데이터 행 스타일
        for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
            for cell in row:
                cell.font = data_font
                cell.border = thin_border

                # 첫 번째 열은 좌측 정렬
                if cell.column == 1:
                    cell.alignment = Alignment(horizontal="left", vertical="center")
                    val = str(cell.value or "")
                    # 섹션 헤더 스타일
                    if val.startswith("【") or val.startswith("▶"):
                        cell.fill = section_fill
                        cell.font = section_font
                        # 해당 행 전체에 섹션 스타일 적용
                        for c in row:
                            c.fill = section_fill
                else:
                    cell.alignment = Alignment(horizontal="right", vertical="center")

        # 열 너비 자동 조정
        for col_idx in range(1, ws.max_column + 1):
            max_len = 10
            col_letter = get_column_letter(col_idx)
            for row in ws.iter_rows(min_col=col_idx, max_col=col_idx):
                for cell in row:
                    if cell.value:
                        cell_len = len(str(cell.value))
                        # 한글은 2배 너비
                        korean_chars = len(re.findall(r'[가-힣]', str(cell.value)))
                        cell_len += korean_chars
                        max_len = max(max_len, cell_len)
            ws.column_dimensions[col_letter].width = min(max_len + 3, 50)

        # 첫 번째 열은 더 넓게
        ws.column_dimensions["A"].width = max(
            ws.column_dimensions["A"].width, 25
        )

        # 틀 고정 (헤더 행 아래)
        ws.freeze_panes = "B2"
